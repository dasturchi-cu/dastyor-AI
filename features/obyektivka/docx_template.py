"""Generate Obyektivka DOCX — ZIP/XML clone of master template (no layout rebuild)."""

from __future__ import annotations

import base64
import logging
import os
import tempfile
from pathlib import Path
from typing import Any

from backend.services.document_render.photo import process_passport_photo
from config.paths import temp_dir
from config.settings import TEMPLATES_DIR
from features.obyektivka.docx_zip import count_page_breaks, read_parts, render_template, write_parts
from features.obyektivka.placeholders import build_placeholder_context

logger = logging.getLogger(__name__)

MASTER_TEMPLATE = TEMPLATES_DIR / "obyektivka_master.docx"


def _to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _decode_photo_data(data: dict[str, Any]) -> str | None:
    photo_data = _to_text(data.get("photo_data") or data.get("img"))
    if not photo_data.startswith("data:image/") or "," not in photo_data:
        return None
    try:
        photo_data = process_passport_photo(photo_data)
        header, b64 = photo_data.split(",", 1)
        mime = header.split(";")[0].split(":")[1].lower()
        ext = {"image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/webp": "webp"}.get(mime, "jpg")
        os.makedirs(str(temp_dir()), exist_ok=True)
        path = str(temp_dir() / f"oby_tpl_photo_{os.getpid()}.{ext}")
        with open(path, "wb") as fh:
            fh.write(base64.b64decode(b64))
        return path
    except Exception as exc:
        logger.warning("photo_data decode failed: %s", exc)
        return None


def _inject_photo_zip(output_path: str, photo_path: str) -> None:
    """Inject photo at reference VML coordinates (page-relative)."""
    if not photo_path or not os.path.isfile(photo_path):
        return
    before = count_page_breaks(Path(output_path))
    try:
        from docx import Document

        from features.obyektivka.docx_picture import add_reference_photo, find_photo_paragraph

        doc = Document(output_path)
        target = find_photo_paragraph(doc)
        if target is None:
            return
        add_reference_photo(target, photo_path)
        for p in doc.paragraphs:
            for run in p.runs:
                if "{{photo}}" in (run.text or ""):
                    run.text = run.text.replace("{{photo}}", "")
        doc.save(output_path)
        after = count_page_breaks(Path(output_path))
        if before != after:
            logger.warning("photo inject changed page breaks (%s -> %s)", before, after)
    except Exception as exc:
        logger.warning("photo inject skipped: %s", exc)


def generate_obyektivka_docx(
    user_data: dict[str, Any] | None = None,
    photo_path: str | None = None,
    output_filepath: str | None = None,
    *,
    watermark: bool = False,
    watermark_text: str | None = None,
    **kwargs: Any,
) -> str:
    data = user_data or kwargs.get("data") or {}
    if not isinstance(data, dict):
        raise ValueError("user_data must be a dictionary")

    if not MASTER_TEMPLATE.is_file():
        raise FileNotFoundError(
            f"Master template missing: {MASTER_TEMPLATE}. Run scripts/build_obyektivka_master.py"
        )

    temp_photo: str | None = None
    if not photo_path or not os.path.exists(photo_path):
        temp_photo = _decode_photo_data(data)
    resolved_photo = photo_path if photo_path and os.path.exists(photo_path) else temp_photo

    if not output_filepath:
        os.makedirs(str(temp_dir()), exist_ok=True)
        safe = (_to_text(data.get("fullname")) or "Obyektivka").replace(" ", "_").replace("/", "_")
        output_filepath = str(temp_dir() / f"obyektivka_{safe}_{os.getpid()}.docx")
    else:
        os.makedirs(os.path.dirname(output_filepath) or ".", exist_ok=True)

    ctx = build_placeholder_context(data)
    render_template(MASTER_TEMPLATE, ctx, Path(output_filepath))

    if resolved_photo:
        _inject_photo_zip(output_filepath, resolved_photo)
    else:
        parts = read_parts(Path(output_filepath))
        xml = parts["word/document.xml"].decode("utf-8")
        if "{{photo}}" in xml:
            xml = xml.replace("{{photo}}", "")
            parts["word/document.xml"] = xml.encode("utf-8")
            write_parts(Path(output_filepath), parts)

    if temp_photo:
        try:
            os.remove(temp_photo)
        except OSError:
            pass

    if watermark:
        from features.obyektivka.docx_watermark import apply_demo_watermark

        apply_demo_watermark(Path(output_filepath), text=watermark_text)

    return output_filepath


def generate_obyektivka_docx_bytes(
    data: dict[str, Any],
    *,
    photo_path: str | None = None,
    watermark: bool = False,
    watermark_text: str | None = None,
) -> bytes:
    """DOCX bytes — master shablon; preview/demo/paid bitta pipeline."""
    path = generate_obyektivka_docx(
        data,
        photo_path=photo_path,
        watermark=watermark,
        watermark_text=watermark_text,
    )
    try:
        return Path(path).read_bytes()
    finally:
        try:
            Path(path).unlink(missing_ok=True)
        except OSError:
            pass
